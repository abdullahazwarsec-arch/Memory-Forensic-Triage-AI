"""
Builds Memory_Forensics_Commands_And_Execution_Runbook.docx — the command
reference for setting up, running, verifying, retraining and publishing
the HazMem project.

Run from the project root:
    python scripts/docs/create_commands_runbook_docx.py
"""

from pathlib import Path

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parents[2]
DOCX_PATH = BASE_DIR / "docs" / "Memory_Forensics_Commands_And_Execution_Runbook.docx"
PROJECT_DIR = r"C:\Users\user\Desktop\Final resarch\Memory_Forensics_Project"
REPO_URL = "https://github.com/abdullahazwarsec-arch/Memory-Forensic-Triage-AI.git"

NAVY = RGBColor(0x1F, 0x38, 0x64)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)
CMD_RED = RGBColor(0x99, 0x1B, 0x1B)


# ── Low-level formatting helpers ──────────────────────────────────────────────
def set_cell_background(cell, fill_hex):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="{fill_hex}"/>'))


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    ))


def set_left_bar(cell, color_hex):
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{color_hex}"/>'
        f'<w:top w:val="nil"/><w:right w:val="nil"/><w:bottom w:val="nil"/>'
        f'</w:tcBorders>'
    ))


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HazMem — Commands & Execution Runbook · Page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    field_run = p.add_run()
    field_run.font.size = Pt(8.5)
    field_run.font.color.rgb = GREY
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        field_run._r.append(el)


# ── Content blocks ────────────────────────────────────────────────────────────
def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    return h


def para(doc, text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    p.paragraph_format.space_after = Pt(6)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):          # (bold lead-in, rest)
            r = p.add_run(item[0])
            r.bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def command(doc, title, cmd, explanation="", output=""):
    """Grey command box: title, the exact command, what it does and what you should see."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F2F5F9")
    set_cell_margins(cell, top=90, bottom=90, left=160, right=160)
    set_left_bar(cell, "1F3864")

    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY

    for line in cmd.split("\n"):
        pc = cell.add_paragraph()
        pc.paragraph_format.space_after = Pt(0)
        rc = pc.add_run(line)
        rc.font.name = "Consolas"
        rc.font.size = Pt(10)
        rc.bold = True
        rc.font.color.rgb = CMD_RED

    if explanation:
        pe = cell.add_paragraph()
        pe.paragraph_format.space_before = Pt(4)
        re_ = pe.add_run(explanation)
        re_.font.size = Pt(9.5)
        re_.font.color.rgb = INK
    if output:
        po = cell.add_paragraph()
        ro = po.add_run("You should see: ")
        ro.bold = True
        ro.font.size = Pt(9.5)
        ro.font.color.rgb = NAVY
        ro2 = po.add_run(output)
        ro2.font.size = Pt(9.5)
        ro2.font.color.rgb = GREY
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def callout(doc, label, text, fill="FFF4E5", bar="C2410C"):
    """Coloured note box for warnings and tips."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill)
    set_cell_margins(cell, top=90, bottom=90, left=160, right=160)
    set_left_bar(cell, bar)
    p = cell.paragraphs[0]
    r = p.add_run(f"{label}  ")
    r.bold = True
    r.font.size = Pt(10)
    p.add_run(text).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def warning(doc, text):
    callout(doc, "⚠ Warning:", text, fill="FFF4E5", bar="C2410C")


def tip(doc, text):
    callout(doc, "💡 Tip:", text, fill="EAF6EE", bar="15803D")


def table(doc, headers, widths, rows, mono_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, name in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        set_cell_background(c, "1F3864")
        set_cell_margins(c, top=80, bottom=80, left=110, right=110)
        r = c.paragraphs[0].add_run(name)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
    for idx, row_data in enumerate(rows):
        cells = t.add_row().cells
        for j, text in enumerate(row_data):
            c = cells[j]
            c.text = ""
            set_cell_background(c, "F7F9FC" if idx % 2 == 0 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=110, right=110)
            for k, line in enumerate(str(text).split("\n")):
                pp = c.paragraphs[0] if k == 0 else c.add_paragraph()
                pp.paragraph_format.space_after = Pt(0)
                rr = pp.add_run(line)
                rr.font.size = Pt(9)
                rr.font.color.rgb = INK
                if j in mono_cols:
                    rr.font.name = "Consolas"
                    rr.font.color.rgb = CMD_RED
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = w
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ── Document ──────────────────────────────────────────────────────────────────
def build_commands_doc():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(0.9)
        add_page_number_footer(s)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    for lvl, size in ((1, 16), (2, 13)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = NAVY

    # Title block
    p = doc.add_paragraph()
    r = p.add_run("HazMem — Memory Forensics AI Triage")
    r.font.size = Pt(22)
    r.bold = True
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    r = p.add_run("Commands & Execution Runbook")
    r.font.size = Pt(15)
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    r = p.add_run("Every command you need to set up, run, verify, retrain and publish the project — "
                  "with what each one does and what you should see.")
    r.italic = True
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(10)

    table(doc, ["Item", "Value"], [Inches(1.8), Inches(4.9)], [
        ("Project folder", PROJECT_DIR),
        ("Shells", "Windows Command Prompt (CMD) and PowerShell"),
        ("Python", "3.11, inside the project virtual environment (venv)"),
        ("Dashboard address", "http://localhost:8501"),
        ("GitHub repository", REPO_URL),
    ], mono_cols=())

    heading(doc, "How to read this guide", 2)
    bullets(doc, [
        ("Grey boxes ", "contain the exact command to type. Red text is the command itself."),
        ("<angle brackets> ", "mean “replace with your own value”, e.g. <dump> → C:\\Evidence\\host01.raw. "
                              "Keep the quotes if the path has spaces."),
        ("Always run commands from the project folder ", "(section 2, step 1) so that data\\, models\\ and outputs\\ are found."),
        ("“python” assumes the virtual environment is active. ", "If it is not, use .\\venv\\Scripts\\python.exe instead."),
    ])

    # 1 ── Quick reference
    heading(doc, "1. Quick Reference — The Commands You Will Use Most", 1)
    table(doc, ["Task", "Command", "Notes"], [Inches(1.7), Inches(3.0), Inches(2.0)], [
        ("Go to the project", f'cd "{PROJECT_DIR}"', "Do this first in every new terminal"),
        ("Activate venv (CMD)", r"venv\Scripts\activate.bat", "Prompt shows (venv)"),
        ("Activate venv (PowerShell)", r".\venv\Scripts\Activate.ps1", "See 2.4 if blocked"),
        ("Start the dashboard", "run_app.bat", "Or: streamlit run app.py"),
        ("Dashboard on this PC only", "streamlit run app.py --server.address localhost", "Recommended with real evidence"),
        ("Stop the dashboard", "Ctrl + C", "In the terminal running it"),
        ("Dump → feature CSV", 'python -m hazmem.memdump_extractor "<dump>"', "Prints before/after hashes"),
        ("Identify a dump's OS", 'vol -f "<dump>" windows.info', "Volatility 3"),
        ("SHA-256 of a file (CMD)", 'certutil -hashfile "<file>" SHA256', "Use MD5 for the MD5 hash"),
        ("SHA-256 of a file (PowerShell)", 'Get-FileHash "<file>" -Algorithm SHA256', ""),
        ("Watch the audit log live", r"Get-Content logs\triage_audit.log -Tail 20 -Wait", "PowerShell; Ctrl + C to stop"),
        ("Check the dataset", "python scripts/evaluation/verify_data_integrity.py", "Rows, duplicates, features"),
        ("Train the model", "python scripts/training/train_model.py", "⚠ Overwrites the dashboard model"),
        ("Evaluate models", "python scripts/evaluation/evaluate_model.py", "⚠ Also overwrites the model"),
        ("Re-pin model hashes", "python scripts/admin/pin_model_hashes.py", "After any re-training"),
        ("Verify the audit log", "python -m hazmem.audit_log verify", "VERIFIED = untouched"),
        ("Unlock / reset an account", "python scripts/admin/manage_users.py --help", "Account recovery"),
        ("Run the automated tests", "python -m pytest", "Needs requirements-dev.txt"),
        ("See what changed (Git)", "git status", ""),
        ("Save & upload to GitHub", 'git add -A\ngit commit -m "<message>"\ngit push origin main', "Three commands, in order"),
        ("Rebuild this document", "python scripts/docs/create_commands_runbook_docx.py", ""),
    ], mono_cols=(1,))

    # 2 ── Setup
    heading(doc, "2. Setting Up the Environment", 1)
    heading(doc, "2.1 Open the project folder", 2)
    command(doc, "Step 1 — Navigate to the project root", f'cd "{PROJECT_DIR}"',
            "Every other command assumes you are in this folder. The quotes are required because the path contains spaces.")
    command(doc, "Check the Python version", "python --version",
            "The project was built and tested with Python 3.11.", "Python 3.11.9")

    heading(doc, "2.2 Create the virtual environment (first time only)", 2)
    command(doc, "Create venv", "python -m venv venv",
            "Creates an isolated Python installation in the venv\\ folder so project packages don't affect the rest of your PC. "
            "Skip this if the venv\\ folder already exists.")

    heading(doc, "2.3 Activate the virtual environment (every new terminal)", 2)
    command(doc, "Command Prompt (CMD)", r"venv\Scripts\activate.bat", "", "(venv) appears at the start of the prompt.")
    command(doc, "PowerShell", r".\venv\Scripts\Activate.ps1", "", "(venv) appears at the start of the prompt.")
    command(doc, "Leave the virtual environment", "deactivate")
    command(doc, "Alternative — no activation needed", r".\venv\Scripts\python.exe <script>.py",
            "Runs the venv's Python directly. Useful in scripts or if activation fails.")

    heading(doc, "2.4 If PowerShell blocks activation", 2)
    command(doc, "Allow scripts for this terminal only", "Set-ExecutionPolicy -Scope Process -ExecutionPolicy Bypass",
            "Only affects the current PowerShell window; it resets when the window closes. Then run Activate.ps1 again.")

    heading(doc, "2.5 Install and check the packages", 2)
    command(doc, "Install everything the project needs", "pip install -r requirements.txt",
            "Installs Streamlit, scikit-learn, SHAP, LIME, ReportLab, Volatility 3 and the rest. Run once, or after requirements.txt changes.")
    command(doc, "Confirm the key packages import",
            'python -c "import streamlit, sklearn, shap, lime, reportlab, volatility3; print(\'OK\')"', "", "OK")
    command(doc, "Show an installed package's version", "pip show volatility3",
            "Replace volatility3 with any package name (streamlit, scikit-learn, shap…).")
    command(doc, "Record exact versions (for your report / reproducibility)", "pip freeze > package_versions.txt",
            "Saves every installed package and version to a text file.")

    heading(doc, "2.6 Files the project needs that are not in Git", 2)
    bullets(doc, [
        ("data\\Obfuscated-MalMem2022.csv ", "— the CIC-MalMem2022 dataset (too large for GitHub). Needed for training, "
                                             "evaluation and the LIME background data. Download it and place it in data\\."),
        ("venv\\ ", "— recreate it with sections 2.2–2.5 on a new computer."),
    ])

    # 3 ── Dashboard
    heading(doc, "3. Running the Dashboard", 1)
    command(doc, "Easiest — double-click or run the launcher", "run_app.bat",
            "Checks the venv exists and starts the dashboard.",
            "Local URL: http://localhost:8501 — the browser opens automatically.")
    command(doc, "Standard command (venv active)", "streamlit run app.py")
    command(doc, "Without activating the venv", r".\venv\Scripts\streamlit.exe run app.py")
    command(doc, "Also works (Run button in VS Code)", "python app.py",
            "app.py starts Streamlit itself when run directly.")
    command(doc, "Use a different port", "streamlit run app.py --server.port 8502",
            "Use when 8501 is already taken. Then open http://localhost:8502.")
    command(doc, "Restrict the dashboard to this computer", "streamlit run app.py --server.address localhost",
            "By default Streamlit also listens on your network (it prints Network/External URLs), so other people "
            "could open it. Use this whenever you handle real evidence.")
    command(doc, "Stop the dashboard", "Ctrl + C", "Press in the terminal window that is running Streamlit.")
    tip(doc, "After editing any .py file (app.py or anything in hazmem\\), "
             "stop the dashboard with Ctrl + C and start it again so every module reloads.")

    heading(doc, "3.1 Using the dashboard (no commands — the clicks)", 2)
    bullets(doc, [
        ("0. Sign in: ", "the first start asks you to create the administrator; administrators add analysts on "
                         "the Users page (temporary password, changed at first sign-in)."),
        ("1. Start a case: ", "enter the case number and the recipient (more details are optional), then Start case. "
                              "The analyst is your account."),
        ("2. Add evidence: ", "say where it came from (device, who collected it, optional expected hash), then choose "
                              "Try a sample, Feature file (CSV) or Memory dump (upload, or a path for large dumps)."),
        ("3. Analyse: ", "press Analyse evidence (pick a snapshot first if the file has several; Scan all snapshots "
                         "scores them all)."),
        ("4. Result: ", "What it means (plain-English verdict, next steps, main reasons) and Report. Switch on "
                        "Show technical details in the sidebar for Explainability, Provenance, Chain of custody and "
                        "Audit log tabs."),
        ("5. Report: ", "download the sealed PDF from the Report tab; its SHA-256 is shown underneath."),
        ("6. Finish: ", "press Close case (confirm) — the closing time is written to the case record and audit log. "
                        "Reopen it later from the Cases page."),
        ("Help: ", "the Help page has a guide, glossary and FAQ; the Assistant button (bottom-right) answers "
                   "questions offline, including “explain my result”."),
    ])

    heading(doc, "3.2 Dashboard settings file", 2)
    para(doc, ".streamlit\\config.toml controls the theme, the maximum upload size (maxUploadSize = 8192 MB) and the toolbar. "
              "Restart the dashboard after changing it.")

    # 4 ── Memory dumps
    heading(doc, "4. Memory Dump Analysis (Volatility 3)", 1)
    heading(doc, "4.1 Convert a dump to the 55-feature CSV", 2)
    command(doc, "Extract features from a memory dump", r'python -m hazmem.memdump_extractor "C:\Evidence\host01_memory.raw"',
            "Hashes the dump, runs windows.info plus 9 Volatility plugins, builds the 55 features, re-hashes the dump "
            "to prove it was not changed, and saves the CSV with a .sha256 sidecar.",
            "progress lines, then “Dump SHA-256 (before)”, “Dump SHA-256 (after) … -> UNCHANGED”, the CSV SHA-256 "
            "and the 55 feature values.")
    command(doc, "Choose the output folder", r'python -m hazmem.memdump_extractor "<dump>" --out "<folder>"',
            "Default output folder: outputs\\extracted\\<dump name>_<date>_<time>\\")
    table(doc, ["File created", "What it is"], [Inches(2.4), Inches(4.3)], [
        ("<dump>_features.csv", "The 55 features — load this in the dashboard as a Feature file (CSV)"),
        ("<dump>_features.csv.sha256", "Hash of the CSV in sha256sum format"),
        ("extraction_report.json", "OS info, plugin timings/row counts, warnings and the before/after hashes"),
        ("windows.<plugin>.json", "Raw output of each Volatility plugin (evidence for review)"),
    ], mono_cols=(0,))
    para(doc, "Supported formats: .raw .mem .vmem .dmp .bin .img .lime .dump — Windows memory images only.")
    warning(doc, "The first analysis of a new Windows version needs internet access: Volatility downloads the kernel "
                 "symbol tables once and caches them. Large dumps take several minutes.")

    heading(doc, "4.2 Running Volatility directly", 2)
    para(doc, "Useful for deeper investigation of a dump after triage. vol is installed in the venv (venv\\Scripts\\vol.exe).")
    table(doc, ["Purpose", "Command"], [Inches(2.3), Inches(4.4)], [
        ("Identify the OS and check the dump is readable", 'vol -f "<dump>" windows.info'),
        ("List running processes", 'vol -f "<dump>" windows.pslist'),
        ("Process tree (parent → child)", 'vol -f "<dump>" windows.pstree'),
        ("Command line of each process", 'vol -f "<dump>" windows.cmdline'),
        ("Injected / suspicious code regions", 'vol -f "<dump>" windows.malfind'),
        ("Hidden processes (cross-view)", 'vol -f "<dump>" windows.psxview'),
        ("Network connections", 'vol -f "<dump>" windows.netscan'),
        ("Services", 'vol -f "<dump>" windows.svcscan'),
        ("Loaded DLLs of one process", 'vol -f "<dump>" windows.dlllist --pid <PID>'),
        ("Save output as JSON", 'vol -q -r json -f "<dump>" windows.pslist > pslist.json'),
        ("List all plugins / help", "vol --help"),
        ("Help for one plugin", 'vol -f "<dump>" windows.malfind --help'),
    ], mono_cols=(1,))

    # 5 ── Integrity
    heading(doc, "5. Evidence Integrity & Chain of Custody", 1)
    para(doc, "The dashboard hashes evidence automatically. These commands let you prove the hashes independently — "
              "for example in front of an examiner.")
    heading(doc, "5.1 Calculate a file's hash", 2)
    command(doc, "SHA-256 (Command Prompt)", 'certutil -hashfile "<file>" SHA256',
            "", "“SHA256 hash of <file>:” followed by 64 hex characters.")
    command(doc, "MD5 (Command Prompt)", 'certutil -hashfile "<file>" MD5')
    command(doc, "SHA-256 (PowerShell)", 'Get-FileHash "<file>" -Algorithm SHA256',
            "Shows the hash in upper case — hashes are not case-sensitive, so it still matches the lower-case values in the ledger.")

    heading(doc, "5.2 Check a file against its .sha256 sidecar", 2)
    command(doc, "Show the recorded hash", 'type "<file>.sha256"', "Format: <hash>  <file name>")
    command(doc, "Compare automatically (PowerShell)",
            '$f = "<file>"\n(Get-FileHash $f -Algorithm SHA256).Hash -eq (Get-Content "$f.sha256").Split(" ")[0]',
            "Prints True if the file is unchanged, False if it was modified.", "True")

    heading(doc, "5.3 Where case records live", 2)
    table(doc, ["Location", "Contents"], [Inches(3.3), Inches(3.4)], [
        (r"outputs\cases\<CASE>\case.json", "Case sessions (analyst, opened/closed times) and evidence items"),
        (r"outputs\cases\<CASE>\<EV-ID>\custody_manifest.json", "Full custody ledger for one evidence item"),
        ("outputs\\cases\\<CASE>\\<EV-ID>\\evidence\\", "Write-protected working copy of uploaded evidence"),
        ("outputs\\cases\\<CASE>\\<EV-ID>\\extraction\\", "Volatility output and feature CSV (memory dumps)"),
        (r"outputs\forensic_report_<ID>_<time>.pdf", "Sealed PDF report (+ .pdf.sha256 sidecar)"),
        (r"logs\triage_audit.log", "HMAC-chained audit log: sign-ins, sessions, custody events, verdicts"),
        ("logs\\archive\\", "Archived audit logs with their verification reports"),
        ("keys\\", "Ledger HMAC key and the account store — back up securely, never share"),
    ], mono_cols=(0,))
    command(doc, "List all cases", r"dir outputs\cases")
    command(doc, "Show a case record", r"type outputs\cases\<CASE>\case.json")

    heading(doc, "5.4 Reading the audit log", 2)
    command(doc, "Show the whole log (CMD)", r"type logs\triage_audit.log")
    command(doc, "Last 30 lines (PowerShell)", r"Get-Content logs\triage_audit.log -Tail 30")
    command(doc, "Follow new entries live (PowerShell)", r"Get-Content logs\triage_audit.log -Tail 20 -Wait",
            "Keep it open while using the dashboard. Ctrl + C to stop.")
    command(doc, "Entries for one case (CMD)", r'findstr "CASE=<CASE>" logs\triage_audit.log')
    command(doc, "Find integrity problems (PowerShell)",
            r'Select-String -Path logs\triage_audit.log -Pattern "MISMATCH|BLOCKED"',
            "No output means no integrity violations were logged.")

    # 6 ── ML scripts
    heading(doc, "6. Machine Learning & Explainability Scripts", 1)
    para(doc, "All scripts are run from the project root as: python scripts/<folder>/<name>.py. "
              "Most need data\\Obfuscated-MalMem2022.csv (see 2.6).")
    warning(doc, "train_model.py and evaluate_model.py both OVERWRITE models\\random_forest_model.pkl — the model the "
                 "dashboard uses. Back it up first (command below). There is no script in the project that re-creates "
                 "models\\isolation_forest.pkl, so never delete that file.")
    command(doc, "Back up the current model before retraining",
            r"copy models\random_forest_model.pkl models\random_forest_model.backup.pkl")
    command(doc, "Restore the backup",
            r"copy /Y models\random_forest_model.backup.pkl models\random_forest_model.pkl")

    table(doc, ["Script (python scripts/…)", "What it does", "Creates / changes"],
          [Inches(2.0), Inches(2.6), Inches(2.1)], [
        ("training/train_model.py", "Trains the binary Random Forest (100 trees) on a random 80/20 split and prints test accuracy",
         "⚠ models/random_forest_model.pkl"),
        ("training/train_multiclass.py", "Trains a Benign / Spyware / Ransomware / Trojan model (not used by the dashboard)",
         "models/multiclass_random_forest.pkl"),
        ("evaluation/evaluate_model.py", "Cross-validation, compares models, draws the plots, saves the best model",
         "outputs/model_comparison.csv, confusion_matrix.png, roc_curve.png\n⚠ models/random_forest_model.pkl"),
        ("evaluation/benchmark.py", "5-fold cross-validation of Random Forest, Gradient Boosting and Logistic Regression",
         "Console output only"),
        ("evaluation/explain_model.py", "Global SHAP summary plot of the trained model", "outputs/shap_summary.png"),
        ("evaluation/lime_example.py", "Prints the top LIME rules for one sample", "Console output only"),
        ("evaluation/verify_data_integrity.py", "Dataset audit: total rows, duplicate rows, label columns, feature count",
         "Console output only"),
        ("admin/manage_users.py", "Account recovery: list, create, reset-password, unlock, enable (audited)",
         "keys/users.json"),
        ("admin/pin_model_hashes.py", "Records the SHA-256 of every model; the dashboard refuses unpinned or modified models",
         "models/model_hashes.json (commit it after re-training)"),
        ("data/check_class_balance.py", "Quick check that the dataset loads and shows the class balance", "Console output only"),
        ("data/generate_synthetic_threats.py", "Creates the rootkit, ransomware and cryptominer test profiles",
         "data/test_samples/fresh_*.csv (overwrites)"),
        ("data/simulate_zeroday.py", "Mutates malware rows into a simulated zero-day profile", "data/apt_zero_day.csv"),
        ("data/extract_real_benign.py", "Copies one benign dataset row as a test profile",
         "data/test_samples/real_benign_sample.csv"),
        ("data/extract_confident_samples.py", "Saves a confidently benign and a confidently malicious row",
         "data/test_samples/guaranteed_benign.csv, guaranteed_malware.csv"),
        ("reports/generate_report.py", "Standalone example PDF incident brief", "A PDF in outputs/"),
        ("reports/pdf_layout_demo.py", "Demonstrates the PDF layout with SHAP and LIME tables", "outputs/test_full_xai_report.pdf"),
        ("docs/create_viva_doc.py", "Builds the viva preparation guide", "docs/Memory_Forensics_AI_Triage_Viva_Defense_Guide.docx"),
        ("docs/create_commands_runbook_docx.py", "Builds this document", "docs/Memory_Forensics_Commands_And_Execution_Runbook.docx"),
    ], mono_cols=(0,))

    heading(doc, "6.1 Verified example outputs", 2)
    command(doc, "Dataset audit", "python scripts/evaluation/verify_data_integrity.py", "",
            "Total Rows: 58,596 · Duplicate Rows: 534 · Ground Truth Label Columns: ['Category', 'Class'] · "
            "Clean Independent Features Count: 55")
    command(doc, "Dataset class balance", "python scripts/data/check_class_balance.py", "", "Class 0 (Benign): 29,298 · Class 1 (Malware): 29,298")
    command(doc, "Verify the audit log chain", "python -m hazmem.audit_log verify", "",
            "VERIFIED  …logs\triage_audit.log  (N chained lines, 584 anchored legacy lines)")
    command(doc, "Verify a custody manifest",
            "python -m hazmem.evidence_integrity verify outputs/cases/<case>/<evidence-id>/custody_manifest.json", "",
            "VERIFIED  …custody_manifest.json  (N events)")
    warning(doc, "The ~100% accuracy printed by train_model.py / evaluate_model.py comes from a random split and is "
                 "optimistic: near-identical snapshots of the same sample appear in both training and test data, and "
                 "the model relies heavily on service/driver counts. See the README “Known Limitations”.")

    heading(doc, "6.2 Quick checks of the trained model", 2)
    command(doc, "Score one CSV with the dashboard model",
            'python -c "import joblib,pandas as pd; m=joblib.load(\'models/random_forest_model.pkl\'); '
            'd=pd.read_csv(\'<file>.csv\').reindex(columns=m.feature_names_in_, fill_value=0); '
            'print(m.predict_proba(d)[:,1])"',
            "Prints the malware probability for every row (0 = benign, 1 = malware).")
    command(doc, "Hash the model file (record it with your results)", r"certutil -hashfile models\random_forest_model.pkl SHA256",
            "Lets you prove which model version produced a result.")

    heading(doc, "6.3 Accounts, audit log and automated tests", 2)
    command(doc, "List dashboard accounts", "python scripts/admin/manage_users.py list")
    command(doc, "Unlock an account / reset a password",
            "python scripts/admin/manage_users.py unlock <username>\npython scripts/admin/manage_users.py reset-password <username>",
            "Use when the only administrator is locked out. Every action is written to the audit log.")
    command(doc, "Archive a broken audit log and start a new chain",
            'python -m hazmem.audit_log archive --reason "<why, at least 10 characters>"',
            "The old log is kept unchanged in logs\\archive\\ with a verification report.")
    command(doc, "Install the test tools (once)", "pip install -r requirements-dev.txt")
    command(doc, "Run all automated tests", "python -m pytest",
            "Security, integrity, validation, audit-log, account and assistant tests (about 2 minutes).",
            "… passed")

    # 7 ── Git
    heading(doc, "7. Git & GitHub", 1)
    command(doc, "First time on a computer — set your identity",
            'git config --global user.name "<Your Name>"\ngit config --global user.email "<you@example.com>"')
    command(doc, "Download the project on a new computer", f"git clone {REPO_URL}",
            "Then recreate the venv (section 2) and add the dataset (2.6).")
    table(doc, ["Task", "Command", "What it does"], [Inches(1.7), Inches(2.6), Inches(2.4)], [
        ("See what changed", "git status", "Lists modified, new and deleted files"),
        ("See the exact changes", "git diff", "Line-by-line changes not yet staged (q to exit)"),
        ("Stage all changes", "git add -A", "Marks every change to be committed"),
        ("Stage one file", "git add <file>", ""),
        ("Commit (save a snapshot)", 'git commit -m "<what you changed>"', "Saves staged changes locally"),
        ("Upload to GitHub", "git push origin main", "Sends your commits to GitHub"),
        ("Get the latest from GitHub", "git pull origin main", "Do this before starting work on another PC"),
        ("History", "git log --oneline -10", "Last 10 commits"),
        ("Undo edits to one file", "git restore <file>", "⚠ Discards your unsaved edits to that file"),
        ("Unstage a file", "git restore --staged <file>", "Keeps your edits, removes them from the next commit"),
    ], mono_cols=(1,))
    para(doc, "Not uploaded on purpose (.gitignore): venv\\, the dataset CSV, everything generated in outputs\\ "
              "(reports, cases, extractions), keys\\ (ledger key, accounts), logs\\, __pycache__ and Word ~$ lock files "
              "— evidence, accounts and case data must never be public.")

    # 8 ── Housekeeping
    heading(doc, "8. Housekeeping", 1)
    command(doc, "Find what is using port 8501", "netstat -ano | findstr :8501", "The last column is the process ID (PID).")
    command(doc, "Stop that process", "taskkill /PID <PID> /F")
    command(doc, "Clear Streamlit's cache", "streamlit cache clear")
    command(doc, "Delete a test case folder",
            r'attrib -r "outputs\cases\<CASE>\*" /s /d' + "\n" + r'rmdir /s /q "outputs\cases\<CASE>"',
            "Evidence working copies are write-protected, so remove the read-only flag first.")
    warning(doc, "Only delete test cases. Deleting a real case destroys its evidence copies and custody records.")

    # 9 ── Troubleshooting
    heading(doc, "9. Troubleshooting", 1)
    table(doc, ["Problem", "Cause", "Fix"], [Inches(2.1), Inches(2.1), Inches(2.5)], [
        ("'streamlit' is not recognized", "venv not active", r".\venv\Scripts\python.exe -m streamlit run app.py"),
        ("Running scripts is disabled (PowerShell)", "Execution policy", "Set-ExecutionPolicy -Scope Process -ExecutionPolicy Bypass"),
        ("ModuleNotFoundError", "venv not active or packages missing", "Activate the venv, then pip install -r requirements.txt"),
        ("Port 8501 is already in use", "Dashboard already running", "Use it, stop it (Ctrl + C / section 8) or add --server.port 8502"),
        ("FileNotFoundError: data/Obfuscated-MalMem2022.csv", "Wrong folder or dataset missing", "cd to the project root; add the dataset (2.6)"),
        ("Volatility: “Unable to validate the plugin requirements”", "Not a Windows memory image, or symbols could not be downloaded",
         "Check the file is a Windows dump; connect to the internet for the first run"),
        ("Upload too large / browser slow", "Uploads are held in memory", "Paste the dump's path in the Memory dump tab instead"),
        ("Dashboard shows old behaviour after editing code", "Modules cached in the running server", "Ctrl + C and start the dashboard again"),
        ("“Analysis blocked — integrity verification failed”", "An evidence file's hash changed since acquisition",
         "Do not continue — check the custody ledger and audit log; re-acquire the evidence"),
        ("Signed out unexpectedly", "Page reloaded, 30 minutes idle, or account disabled",
         "Sign in again; use the top menu to move between pages"),
        ("“Invalid username or password” with the right password", "Account locked after 5 failed attempts",
         "Wait 15 minutes, or: python scripts/admin/manage_users.py unlock <username>"),
        ("“Audit log integrity check FAILED”", "A log line was altered, removed or duplicated",
         "Read the diagnosis on the Audit log page; archive the log (section 6.3) — never edit it"),
        ("PermissionError when saving a report or document", "The file is open in a PDF viewer or Word", "Close the file and try again"),
        ("~$…docx files appear", "Word lock files while a document is open", "Harmless — they disappear when Word closes"),
    ], mono_cols=())

    doc.save(str(DOCX_PATH))
    print(f"[SUCCESS] Commands runbook saved to: {DOCX_PATH}")


if __name__ == "__main__":
    build_commands_doc()
